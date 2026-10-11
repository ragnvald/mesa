"""Report layout rules: map views, area figures, the geocode table and Word page flow."""
from __future__ import annotations

import os
import sys
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

ROOT = Path(__file__).resolve().parents[1]
CODE_DIR = ROOT / "code"
if str(CODE_DIR) not in sys.path:
    sys.path.insert(0, str(CODE_DIR))

import geopandas as gpd  # noqa: E402
import pandas as pd  # noqa: E402
from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from PIL import Image  # noqa: E402
from shapely.geometry import box  # noqa: E402

import report_generate as rg  # noqa: E402


def test_a_tall_extent_is_widened_to_the_page_panel() -> None:
    # 100 km wide, 400 km tall: height-bound on the page unless the view is widened.
    view, fig_w, fig_h = rg._page_map_view((0.0, 0.0, 100_000.0, 400_000.0))
    width, height = view[2] - view[0], view[3] - view[1]

    assert abs(height / width - rg.MAP_PANEL_ASPECT) < 1e-6
    assert abs(fig_h / fig_w - rg.MAP_PANEL_ASPECT) < 1e-6
    # the data sits inside the view with a margin on every side
    assert view[0] < 0.0 and view[2] > 100_000.0 and view[1] < 0.0 and view[3] > 400_000.0
    assert (0.0 - view[1]) >= rg.MAP_PAD_RATIO * 400_000.0 - 1e-6


def test_a_wide_extent_keeps_its_width_and_gets_a_margin() -> None:
    view, _w, _h = rg._page_map_view((0.0, 0.0, 400_000.0, 250_000.0))
    assert view[0] < 0.0 and view[2] > 400_000.0
    assert 0.5 <= (view[3] - view[1]) / (view[2] - view[0]) <= rg.MAP_PANEL_ASPECT + 1e-6


def test_area_figures_use_digits_that_match_their_size() -> None:
    assert rg._format_area_km2(4181.8031) == "4,182 km²"
    assert rg._format_area_km2(766.408) == "766.4 km²"
    assert rg._format_area_km2(71.0041) == "71.00 km²"
    assert rg._format_area_km2(3.5102) == "3.510 km²"
    assert rg._format_area_km2(0.25) == "250,000 m²"
    assert rg._format_area_compact(577_000) == "57.7 ha"
    assert rg._format_area_compact(float("nan")) == "–"


def test_geocode_overview_has_average_and_total_area(tmp_path: Path) -> None:
    # Two 0.1 x 0.1 degree cells on the equator (~123 km² each) and one four times the size.
    cells = gpd.GeoDataFrame(
        {"name_gis_geocodegroup": ["grid", "grid", "big"], "code": ["a", "b", "c"]},
        geometry=[box(0, 0, 0.1, 0.1), box(0.1, 0, 0.2, 0.1), box(0, 0, 0.2, 0.2)],
        crs="EPSG:4326",
    )
    cells.to_parquet(tmp_path / "tbl_geocode_object.parquet", index=False)
    pd.DataFrame({"name_gis_geocodegroup": ["grid", "big"],
                  "title_user": ["Grid", float("nan")],
                  "description": [None, "One large cell"]}).to_parquet(tmp_path / "tbl_geocode_group.parquet", index=False)

    rows = rg.build_geocode_overview_table(str(tmp_path))
    assert rows[0] == ["Geocode group", "Title", "Description", "# objects", "Average area", "Total area"]
    by_name = {r[0]: r for r in rows[1:]}

    assert by_name["grid"][3] == "2"
    assert by_name["grid"][4].endswith("km²") and by_name["grid"][5].endswith("km²")
    grid_avg = float(by_name["grid"][4].split()[0].replace(",", ""))
    grid_total = float(by_name["grid"][5].split()[0].replace(",", ""))
    big_avg = float(by_name["big"][4].split()[0].replace(",", ""))
    assert 115 < grid_avg < 130
    assert abs(grid_total - 2 * grid_avg) < 2
    assert abs(big_avg - 4 * grid_avg) < 8
    # missing title / description never print as the word "nan"
    assert by_name["big"][1] == "" and by_name["grid"][2] == ""


def test_study_area_map_fills_the_page_panel(tmp_path: Path, monkeypatch) -> None:
    # The scale bar once got the frame in the wrong coordinate order; its label landed far
    # outside the map and the saved image became a strip five times taller than wide.
    monkeypatch.setattr(rg, "_plot_basemap", lambda *a, **k: None)
    area = gpd.GeoDataFrame({"title": ["a"]}, geometry=[box(-1.0, 8.2, -0.5, 9.6)], crs="EPSG:4326")
    cells = gpd.GeoDataFrame({"sensitivity_code_max": ["A", "C"]},
                             geometry=[box(-3.0, 5.0, 0.0, 12.0), box(0.0, 5.0, 3.0, 12.0)], crs="EPSG:4326")
    out = tmp_path / "area.png"

    assert rg._analysis_write_area_map_png(area, str(out), base_dir=None, config_path=None,
                                           palette_A2E={"A": "#bd0026", "C": "#fd8d3c"}, desc_A2E={},
                                           context_cells_gdf=cells)
    width, height = Image.open(out).size
    assert abs(height / width - rg.MAP_PANEL_ASPECT) < 0.05


def test_word_page_flow(tmp_path: Path) -> None:
    picture = tmp_path / "geocode_demo_map_sensitivity.png"
    legend = tmp_path / "geocode_demo_sensitivity_legend.png"
    Image.new("RGB", (800, 600), "#cccccc").save(picture)
    Image.new("RGB", (800, 60), "#eeeeee").save(legend)

    order = [
        ('heading(2)', "First section"),
        ('text', "Opening text."),
        ('new_page', None),
        ('heading(3)', "Importance (max)"),
        ('text', "What the map shows."),
        ('image_map', ("Importance (max)", str(picture))),
        ('image', (None, str(legend))),
        ('new_page', None),
        ('heading(2)', "A table"),
        ('table_data', ("Counts", [["Name", "Value"], ["a", 1], ["b", 2]])),
    ]
    out = tmp_path / "report.docx"
    rg.compile_docx(str(out), order)
    doc = Document(str(out))

    # A4 with a 16 cm text column, the width the images are sized for
    section = doc.sections[0]
    assert round(section.page_width.cm, 1) == 21.0 and round(section.page_height.cm, 1) == 29.7
    text_width_cm = section.page_width.cm - section.left_margin.cm - section.right_margin.cm
    assert round(text_width_cm, 1) == 16.0

    paragraphs = doc.paragraphs
    texts = [p.text for p in paragraphs]

    # page breaks are a property of the paragraph that starts the page, never a paragraph of their own
    xml = doc.element.xml
    assert 'w:br w:type="page"' not in xml
    starts_page = [p.text for p in paragraphs if p.paragraph_format.page_break_before]
    assert starts_page == ["Importance (max)", "A table"]

    # the bold figure title that only repeats the heading is not printed
    assert texts.count("Importance (max)") == 1

    # heading -> introduction -> map -> attribution -> legend is one block
    intro = next(p for p in paragraphs if p.text == "What the map shows.")
    assert intro.paragraph_format.keep_with_next
    picture_paragraphs = [p for p in paragraphs if p._p.findall('.//' + qn('w:drawing'))]
    assert len(picture_paragraphs) == 2
    assert picture_paragraphs[0].paragraph_format.keep_with_next
    attribution = next(p for p in paragraphs if "OpenStreetMap" in p.text)
    assert attribution.paragraph_format.keep_with_next
    # text that is not followed by a figure is left alone
    assert not next(p for p in paragraphs if p.text == "Opening text.").paragraph_format.keep_with_next

    # no picture is wider than the text column
    for shape in doc.inline_shapes:
        assert shape.width.cm <= 16.01

    # table rows never split across pages and the header row repeats
    table = doc.tables[0]
    for row in table.rows:
        assert row._tr.trPr is not None and row._tr.trPr.find(qn('w:cantSplit')) is not None
    assert table.rows[0]._tr.trPr.find(qn('w:tblHeader')) is not None



def _type_records() -> list:
    cols = [f"h_i{i}_s{s}" for i in range(1, 6) for s in range(1, 6)]
    uniform = dict.fromkeys(cols, 0.0) | {"h_i5_s3": 1.0}
    mixed = dict.fromkeys(cols, 0.0) | {"h_i4_s4": 0.6, "h_i2_s5": 0.4}
    return [
        {"cluster_id": 0, "cluster_label": "type 1", "n_polygons": 1200, "total_area_km2": 990.0,
         "mean_importance": 5.0, "mean_susceptibility": 3.0, "mean_coverage_index": 1.0,
         "top_asset_groups": "settlements,roads", **uniform},
        {"cluster_id": 1, "cluster_label": "type 2", "n_polygons": 40, "total_area_km2": 5.0,
         "mean_importance": 3.2, "mean_susceptibility": 4.4, "mean_coverage_index": 2.5,
         "top_asset_groups": "", **mixed},
    ]


def test_sensitivity_types_are_listed_as_a_table() -> None:
    rows = rg._classification_type_table(_type_records(), 995.0)
    assert rows[0] == ["Type", "Character", "Area", "Cells", "Mean imp. / sus.", "Overlap depth", "Dominant asset groups"]
    assert rows[1] == ["type 1", "High importance, moderate susceptibility", "990.0 km² (99%)", "1,200",
                       "5.0 / 3.0", "1.0", "settlements, roads"]
    # a type too small to round to a whole percent is not shown as 0%
    assert rows[2][2] == "5.000 km² (<1%)" and rows[2][1] == "Moderate importance, high susceptibility"


def test_type_fingerprints_are_drawn_one_grid_per_type(tmp_path: Path) -> None:
    records = _type_records()
    cols = [c for c in records[0] if c.startswith("h_i")]
    out = tmp_path / "fingerprints.png"
    assert rg._type_fingerprints_png(records, cols, str(out), {"type 1": "#8dd3c7"}, 995.0)
    width, height = Image.open(out).size
    assert width > 1.6 * height            # two grids side by side
    assert not rg._type_fingerprints_png([], cols, str(tmp_path / "none.png"))
